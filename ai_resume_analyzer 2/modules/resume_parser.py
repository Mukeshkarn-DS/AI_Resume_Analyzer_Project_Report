"""
Resume Parser Module
Extracts raw text from PDF and DOCX files.
"""

import pdfplumber
from docx import Document
import os


def extract_text(filepath):
    """
    Extract text from PDF or DOCX file.

    Args:
        filepath (str): Path to the resume file.

    Returns:
        str: Extracted text.
    """
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"File not found: {filepath}")

    ext = filepath.lower().rsplit(".", 1)[-1]

    if ext == "pdf":
        return _extract_from_pdf(filepath)
    elif ext == "docx":
        return _extract_from_docx(filepath)
    else:
        raise ValueError(f"Unsupported file format: {ext}")


def _extract_from_pdf(filepath):
    """Extract text from PDF using pdfplumber."""
    text = ""
    try:
        with pdfplumber.open(filepath) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception as e:
        raise RuntimeError(f"PDF parsing error: {str(e)}")

    return text.strip()


def _extract_from_docx(filepath):
    """Extract text from DOCX using python-docx."""
    text_parts = []
    try:
        doc = Document(filepath)

        # Extract from paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)

        # Extract from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip():
                    text_parts.append(row_text)

    except Exception as e:
        raise RuntimeError(f"DOCX parsing error: {str(e)}")

    return "\n".join(text_parts).strip()